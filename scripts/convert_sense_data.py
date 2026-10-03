"""Convert the raw VILPE Sense export (zip of xlsx/docx/pdf) into the
categorized dataset under data/.

Run from the repo root:

    uv run --with openpyxl --with python-docx \
        python scripts/convert_sense_data.py <source.zip|dir> [data]
"""

import csv
import json
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path

import docx
import openpyxl

PURPOSE_EN = {
    "Kattorakenteen tuuletus": "Roof structure ventilation",
    "Ryömintätilan tuuletus": "Crawl space ventilation",
}

FAN_COLUMNS = {
    "Ohjaava sisälähetin T (°C)": "indoor_temp_c",
    "Ohjaava sisälähetin H (%)": "indoor_rh_pct",
    "Ohjaava sisälähetin g/m³": "indoor_abs_humidity_g_m3",
    "Ohjaava ulkolähetin T (°C)": "outdoor_temp_c",
    "Ohjaava ulkolähetin H (%)": "outdoor_rh_pct",
    "Ohjaava ulkolähetin g/m³": "outdoor_abs_humidity_g_m3",
}
FAN_CSV_HEADER = ["timestamp", "rpm", *FAN_COLUMNS.values()]


def slugify(name: str) -> str:
    slug = name.lower().replace("vilpe vantaa,", "").strip()
    slug = re.sub(r"[^a-z0-9]+", "-", slug).strip("-")
    return slug


def location_of(name: str) -> str:
    if "viherkatto" in name.lower():
        return "green_roof"
    if "alapohja" in name.lower():
        return "crawl_space"
    return "roof"


def iso(ts) -> str:
    if isinstance(ts, datetime):
        return ts.isoformat(timespec="seconds")
    return datetime.strptime(str(ts).strip(), "%Y/%m/%d %H:%M:%S").isoformat(
        timespec="seconds"
    )


def parse_fan_workbook(path: Path) -> tuple[dict, list[list]]:
    """Return (device metadata, time-series rows) from one fan export."""
    ws = openpyxl.load_workbook(path, read_only=True)["Sheet1"]
    rows = [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]

    META_KEYS = {
        "Sarjanumero",
        "Tyyppi",
        "Omistaja",
        "Tila",
        "Käyttötarkoitus",
        "Rakennusmateriaali",
    }
    meta = {}
    paired = {}
    series = []
    section = "meta"
    for i, row in enumerate(rows):
        key = row[0]
        if section == "meta":
            if key == "Tyyppi" and row[1] == "Tunniste":
                section = "paired"
            elif key == "Tunniste" and i == 0:
                meta["name"] = row[1]
            elif key in META_KEYS:
                meta[key] = row[1]
            elif key == "Viimeisin homeindeksi":
                meta["mold_index"] = row[1]
            elif key == "Viimeisin rpm":
                meta["rpm"] = row[1]
        elif section == "paired":
            if key in ("Ohjaava sisälähetin", "Ohjaava ulkolähetin"):
                side = "indoor" if "sisä" in key else "outdoor"
                paired[side] = {
                    "serial_number": row[2],
                    "last_seen": iso(row[3]),
                    "temperature_c": row[4],
                    "rh_pct": row[5],
                }
            elif key == "Aikaleima" and row[1] == "Rpm":
                section = "series"
                mapping = {
                    j: FAN_COLUMNS[h] for j, h in enumerate(row) if h in FAN_COLUMNS
                }
        elif section == "series" and key is not None:
            out = {c: None for c in FAN_CSV_HEADER}
            out["timestamp"] = iso(key)
            out["rpm"] = row[1]
            for j, col in mapping.items():
                out[col] = row[j]
            series.append([out[c] for c in FAN_CSV_HEADER])

    name = meta["name"]
    device = {
        "id": slugify(name),
        "name": name,
        "serial_number": meta["Sarjanumero"],
        "type": meta["Tyyppi"],
        "category": "ventilation_fan",
        "location": location_of(name),
        "purpose": {
            "fi": meta.get("Käyttötarkoitus"),
            "en": PURPOSE_EN.get(meta.get("Käyttötarkoitus")),
        },
        "building_material": meta.get("Rakennusmateriaali"),
        "owner": meta.get("Omistaja"),
        "status": meta.get("Tila"),
        "paired_sensors": paired,
        "latest": {"mold_index": meta.get("mold_index"), "rpm": meta.get("rpm")},
    }
    return device, series


def parse_measurements(path: Path):
    ws = openpyxl.load_workbook(path, read_only=True)["Sheet1"]
    it = ws.iter_rows(values_only=True)
    next(it)  # header
    readings, sensors, bases = [], {}, {}
    for r in it:
        if r[0] is None:
            continue
        ts, sid, sserial, _sname, bid, bserial, bname, t, rh, ah = r[:10]
        readings.append(
            [iso(ts), sid, sserial, bid, bserial, t, rh, ah]
        )
        s = sensors.setdefault(
            sserial,
            {
                "sensor_id": sid,
                "serial_number": sserial,
                "reading_count": 0,
                "first_seen": ts,
                "last_seen": ts,
            },
        )
        s["reading_count"] += 1
        s["first_seen"] = min(s["first_seen"], ts)
        s["last_seen"] = max(s["last_seen"], ts)
        bases[bserial] = {"id": bid, "serial_number": bserial, "name": bname}
    return readings, list(sensors.values()), list(bases.values())


def parse_site_doc(path: Path) -> dict:
    d = docx.Document(str(path))
    paragraphs = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    links = {}
    site_link = None
    for table in d.tables:
        for row in table.rows[1:]:
            serial, name, link = (c.text.strip() for c in row.cells[:3])
            if serial:
                links[serial] = link
            elif "site/" in link:
                site_link = link
    return {
        "title": paragraphs[0] if paragraphs else None,
        "description": "\n".join(paragraphs[1:]),
        "device_links": links,
        "site_link": site_link,
    }


def main() -> None:
    src = Path(sys.argv[1])
    outdir = Path(sys.argv[2] if len(sys.argv) > 2 else "data")

    tmp = None
    if src.suffix == ".zip":
        tmp = tempfile.TemporaryDirectory()
        zipfile.ZipFile(src).extractall(tmp.name)
        src = Path(tmp.name)

    (outdir / "readings" / "fans").mkdir(parents=True, exist_ok=True)

    # --- fans ---
    devices = []
    paired_serials = {}
    for f in sorted(src.glob("VILPE*.xlsx")):
        device, series = parse_fan_workbook(f)
        device["readings_file"] = f"readings/fans/{device['id']}.csv"
        device["readings_count"] = len(series)
        devices.append(device)
        for side, s in device["paired_sensors"].items():
            paired_serials[s["serial_number"]] = {"device": device["id"], "side": side}
        with open(outdir / device["readings_file"], "w", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(FAN_CSV_HEADER)
            w.writerows(series)

    # --- sensor measurements ---
    meas_file = next(src.glob("measurements_*.xlsx"))
    readings, sensors, bases = parse_measurements(meas_file)
    for s in sensors:
        s["first_seen"] = iso(s["first_seen"])
        s["last_seen"] = iso(s["last_seen"])
        pair = paired_serials.get(s["serial_number"])
        s["paired_to"] = pair["device"] if pair else None
        s["role"] = "paired_transmitter" if pair else "environment_sensor"
    sensors.sort(key=lambda s: s["serial_number"])
    with open(outdir / "readings" / "sensors.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(
            [
                "timestamp",
                "sensor_id",
                "sensor_serial",
                "base_station_id",
                "base_station_serial",
                "temperature_c",
                "relative_humidity_pct",
                "absolute_humidity_g_m3",
            ]
        )
        w.writerows(readings)

    # --- site doc + layout ---
    doc_file = next(src.glob("*.docx"))
    site = parse_site_doc(doc_file)
    for d in devices:
        d["public_link"] = site["device_links"].get(d["serial_number"])
    site_json = {
        "id": "vilpe-vantaa",
        "name": "VILPE Vantaa",
        "description_fi": site["description"],
        "public_site_link": site["site_link"],
        "timezone": "Europe/Helsinki",
        "base_stations": bases,
        "layout_file": "site-layout.pdf",
    }
    pdf = next(src.glob("site_layout_*.pdf"), None)
    if pdf:
        (outdir / "site-layout.pdf").write_bytes(pdf.read_bytes())

    (outdir / "site.json").write_text(
        json.dumps(site_json, indent=2, ensure_ascii=False) + "\n"
    )
    (outdir / "devices.json").write_text(
        json.dumps(devices, indent=2, ensure_ascii=False) + "\n"
    )
    (outdir / "sensors.json").write_text(
        json.dumps(sensors, indent=2, ensure_ascii=False) + "\n"
    )

    n_paired = sum(1 for s in sensors if s["role"] == "paired_transmitter")
    print(f"devices: {len(devices)}")
    for d in devices:
        print(f"  {d['id']}: {d['readings_count']} rows")
    print(
        f"sensors: {len(sensors)} ({n_paired} paired transmitters), "
        f"{len(readings)} readings"
    )
    print(f"base stations: {bases}")
    if tmp:
        tmp.cleanup()


if __name__ == "__main__":
    main()
